# -*- coding: utf-8 -*-
"""DOCX document parser using python-docx."""

import uuid
from typing import Tuple, List, Dict, Any
import docx
from backend.agents.document_intelligence.parsers.base_parser import BaseParser
from backend.agents.document_intelligence.exceptions import ParserFailureError
from backend.agents.document_intelligence.schema import (
    Page, Table, Cell, FormField, ExtractedTextBlock, SourceReference, DocumentElementType, Paragraph
)

class DOCXParser(BaseParser):
    """Parses Word document DOCX files into structural text and tables blocks."""
    
    def parse(
        self, 
        file_path: str, 
        evidence_id: str
    ) -> Tuple[List[Page], List[Table], List[FormField], List[ExtractedTextBlock]]:
        pages: List[Page] = []
        tables: List[Table] = []
        forms: List[FormField] = []
        raw_blocks: List[ExtractedTextBlock] = []
        
        reading_order_counter = 0
        
        try:
            doc = docx.Document(file_path)
            
            # DOCX file has no physical page breaks stored in headers. We treat it as one continuous page (Page 1)
            page_num = 1
            paragraphs_list: List[Paragraph] = []
            full_text_accumulator = []
            
            # 1. Parse paragraphs
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                    
                full_text_accumulator.append(text)
                
                # Check style for header classification
                source_element = DocumentElementType.PARAGRAPH
                if para.style.name.startswith("Heading") or para.style.name == "Title":
                    source_element = DocumentElementType.TITLE
                    
                ref = SourceReference(
                    evidence_id=evidence_id,
                    file_path=file_path,
                    page_number=page_num,
                    parser="DOCXParser",
                    parser_version="1.0.0",
                    source_element=source_element
                )
                
                raw_blocks.append(
                    ExtractedTextBlock(
                        text=text,
                        reading_order_index=reading_order_counter,
                        source_reference=ref
                    )
                )
                
                paragraphs_list.append(
                    Paragraph(
                        paragraph_id=str(uuid.uuid4()),
                        page_number=page_num,
                        text=text,
                        reading_order=reading_order_counter
                    )
                )
                reading_order_counter += 1
                
            # 2. Parse Word tables
            page_tables: List[Table] = []
            for doc_table in doc.tables:
                table_id = str(uuid.uuid4())
                cells: List[Cell] = []
                
                rows_count = len(doc_table.rows)
                cols_count = len(doc_table.columns) if rows_count > 0 else 0
                
                headers: List[str] = []
                
                for r_idx, row in enumerate(doc_table.rows):
                    row_headers = []
                    for c_idx, cell in enumerate(row.cells):
                        cell_text = cell.text.strip()
                        cell_ref = SourceReference(
                            evidence_id=evidence_id,
                            file_path=file_path,
                            page_number=page_num,
                            parser="DOCXParser",
                            parser_version="1.0.0",
                            source_element=DocumentElementType.TABLE_CELL
                        )
                        cells.append(
                            Cell(
                                text=cell_text,
                                row_index=r_idx,
                                col_index=c_idx,
                                source_reference=cell_ref
                            )
                        )
                        if r_idx == 0:
                            row_headers.append(cell_text)
                    if r_idx == 0:
                        headers = row_headers
                        
                table_obj = Table(
                    table_id=table_id,
                    page_number=page_num,
                    rows_count=rows_count,
                    cols_count=cols_count,
                    cells=cells,
                    headers=headers,
                    reading_order=reading_order_counter
                )
                tables.append(table_obj)
                page_tables.append(table_obj)
                reading_order_counter += 1
                
            pages.append(
                Page(
                    page_number=page_num,
                    text_content="\n\n".join(full_text_accumulator),
                    paragraphs=paragraphs_list,
                    tables=page_tables,
                    forms=[]
                )
            )
            
            return pages, tables, forms, raw_blocks
            
        except Exception as e:
            raise ParserFailureError(f"Failed parsing DOCX file: {e}")
